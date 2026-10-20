from io import BytesIO

from openpyxl import load_workbook


def extract_queries(data: bytes) -> list[str]:
    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.worksheets[0]

        values = [
            row[0]
            for row in sheet.iter_rows(
                min_col=1,
                max_col=1,
                values_only=True,
            )
        ]
    finally:
        workbook.close()

    return [
        str(int(value) if isinstance(value, float) and value.is_integer() else value).strip()
        for value in values
        if value is not None and str(value).strip()
    ]