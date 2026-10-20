from io import BytesIO

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.models import Status, Store, Job


def extract_queries(data: bytes) -> list[str]:
    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.worksheets[0]

        values = [
            row[0]
            for row in sheet.iter_rows(
                min_col=1,
                max_col=1,
                values_only=True,
            )
        ]
    finally:
        workbook.close()

    return [
        str(int(value) if isinstance(value, float) and value.is_integer() else value).strip()
        for value in values
        if value is not None and str(value).strip()
    ]

def export_job_excel(job: Job) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сравнение цен"

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )
    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE8E6",
    )

    stores = [
        Store.dns,
        Store.citilink,
        Store.onlinetrade,
    ]

    sheet.append([
        "Артикул",
        "DNS",
        "Citilink",
        "Onlinetrade",
    ])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    status_text = {
        Status.not_found: "Не найдено",
        Status.failed: "Ошибка",
        Status.canceled: "Отменено",
    }

    for product in sorted(job.products, key=lambda p: p.position):
        searches = {
            search.store: search
            for search in product.searches
        }

        row = sheet.max_row + 1
        sheet.cell(row=row, column=1, value=product.query)

        prices = {}

        for column, store in enumerate(stores, start=2):
            search = searches.get(store)
            if search is None:
                continue

            cell = sheet.cell(row=row, column=column)

            if search.status == Status.found:
                if search.offers:
                    price = min(
                        offer.price
                        for offer in search.offers
                    )

                    prices[column] = price
                    cell.value = float(price)
                    cell.number_format = '#,##0.00'

            elif search.status in status_text:
                cell.value = status_text[search.status]
                cell.fill = red_fill

            # pending и in_progress остаются пустыми

        if prices:
            lowest_price = min(prices.values())

            for column, price in prices.items():
                if price == lowest_price:
                    sheet.cell(
                        row=row,
                        column=column,
                    ).fill = green_fill

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 20

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    return output.getvalue()